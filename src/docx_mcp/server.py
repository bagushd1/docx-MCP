import os
from pathlib import Path
from mcp.server.mcpserver import MCPServer
from docx import Document
from docx.shared import Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH

mcp = MCPServer("DocxEditor")


def validate_docx_path(file_path: str, must_exist: bool = False) -> Path:
    if not file_path or not file_path.strip():
        raise ValueError("Path file tidak boleh kosong.")

    resolved_path = Path(file_path).expanduser().resolve()

    if resolved_path.suffix.lower() != ".docx":
        raise ValueError(
            f"Operasi ditolak: '{resolved_path.name}' bukan file .docx yang valid."
        )

    allowed_dir = os.environ.get("DOCX_MCP_ALLOWED_DIR")
    if allowed_dir:
        allowed_base = Path(allowed_dir).expanduser().resolve()
        if not resolved_path.is_relative_to(allowed_base):
            raise PermissionError(
                f"Akses ditolak: File berada di luar direktori yang diizinkan ({allowed_base})."
            )

    if must_exist and not resolved_path.is_file():
        raise FileNotFoundError(f"File '{resolved_path}' tidak ditemukan.")

    return resolved_path


@mcp.tool()
def read_docx(file_path: str) -> str:
    """Membaca seluruh isi teks dan tabel dari file .docx"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))
    except Exception as e:
        return f"Error: Gagal membaca file {file_path}: {e}"

    content = []

    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            content.append(f"[P{i}] {p.text}")

    for t_idx, table in enumerate(doc.tables):
        content.append(f"\n--- Tabel {t_idx} ---")
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            content.append(" | ".join(row_text))

    return "\n".join(content)


@mcp.tool()
def create_or_append_docx(file_path: str, paragraphs: list[str], heading: str = "") -> str:
    """Membuat file .docx baru atau menambahkan paragraf ke file yang sudah ada"""
    try:
        path = validate_docx_path(file_path, must_exist=False)
        doc = Document(str(path)) if path.is_file() else Document()

        if heading:
            doc.add_heading(heading, level=1)

        for text in paragraphs:
            doc.add_paragraph(text)

        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))
        return f"Berhasil menyimpan perubahan ke {path}"
    except Exception as e:
        return f"Error: Gagal menyimpan file {file_path}: {e}"


@mcp.tool()
def replace_text_in_docx(file_path: str, old_text: str, new_text: str) -> str:
    """Mencari dan mengganti teks spesifik di dalam paragraf maupun tabel .docx"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))
        count = 0

        for p in doc.paragraphs:
            if old_text in p.text:
                count += p.text.count(old_text)
                p.text = p.text.replace(old_text, new_text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if old_text in cell.text:
                        count += cell.text.count(old_text)
                        cell.text = cell.text.replace(old_text, new_text)

        doc.save(str(path))
        return f"Berhasil mengganti {count} kemunculan '{old_text}' menjadi '{new_text}' di {path}."
    except Exception as e:
        return f"Error: Gagal memproses file {file_path}: {e}"


@mcp.tool()
def create_table_docx(file_path: str, data: list[list[str]], heading: str = "") -> str:
    """Membuat tabel baru dari matriks data 2D ke dalam file .docx"""
    try:
        if not data or not any(data):
            return "Error: Data tabel tidak boleh kosong."

        path = validate_docx_path(file_path, must_exist=False)
        doc = Document(str(path)) if path.is_file() else Document()

        if heading:
            doc.add_heading(heading, level=1)

        num_rows = len(data)
        num_cols = max(len(row) for row in data)
        try:
            table = doc.add_table(rows=num_rows, cols=num_cols, style="Table Grid")
        except Exception:
            table = doc.add_table(rows=num_rows, cols=num_cols)

        for r_idx, row_values in enumerate(data):
            for c_idx, val in enumerate(row_values):
                table.rows[r_idx].cells[c_idx].text = str(val)

        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))
        return f"Berhasil membuat tabel ({num_rows} baris, {num_cols} kolom) di {path}."
    except Exception as e:
        return f"Error: Gagal membuat tabel di {file_path}: {e}"


@mcp.tool()
def edit_table_docx(file_path: str, table_index: int, row: int, col: int, new_text: str) -> str:
    """Mengubah isi teks pada sel tabel tertentu (0-indexed) dalam file .docx"""
    # shortcut: single-cell edit only; add batch edit or row/column insertion if requested.
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]
        if row < 0 or row >= len(table.rows):
            return f"Error: Indeks baris {row} tidak valid. Total baris: {len(table.rows)}."

        target_row = table.rows[row]
        if col < 0 or col >= len(target_row.cells):
            return f"Error: Indeks kolom {col} tidak valid. Total kolom pada baris {row}: {len(target_row.cells)}."

        target_row.cells[col].text = new_text
        doc.save(str(path))
        return f"Berhasil mengubah sel [{row}, {col}] pada tabel {table_index} di {path}."
    except Exception as e:
        return f"Error: Gagal mengedit tabel di {file_path}: {e}"


@mcp.tool()
def delete_table_docx(file_path: str, table_index: int) -> str:
    """Menghapus tabel dari file .docx berdasarkan indeks tabel (0-indexed)"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]
        tbl_elem = table._element
        tbl_elem.getparent().remove(tbl_elem)
        doc.save(str(path))
        return f"Berhasil menghapus tabel {table_index} dari {path}."
    except Exception as e:
        return f"Error: Gagal menghapus tabel di {file_path}: {e}"


@mcp.tool()
def clear_docx(file_path: str) -> str:
    """Menghapus seluruh isi teks dan tabel dari file .docx sehingga menjadi dokumen kosong"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document()
        doc.save(str(path))
        return f"Berhasil mengosongkan seluruh isi file {path}."
    except Exception as e:
        return f"Error: Gagal mengosongkan file {file_path}: {e}"


@mcp.tool()
def add_column_docx(
    file_path: str,
    table_index: int,
    values: list[str] | None = None,
    width_inches: float | None = None,
) -> str:
    """Menambahkan satu kolom baru pada tabel docx dengan teks sel opsional dan lebar opsional"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]
        width_val = Inches(width_inches) if width_inches and width_inches > 0 else None
        table.add_column(width=width_val)

        if width_val:
            table.columns[-1].width = width_val

        if values:
            for idx, val in enumerate(values[: len(table.rows)]):
                cell = table.rows[idx].cells[-1]
                cell.text = str(val)
                if width_val:
                    cell.width = width_val
        elif width_val:
            for row in table.rows:
                row.cells[-1].width = width_val

        doc.save(str(path))
        new_col_idx = len(table.columns) - 1
        return f"Berhasil menambahkan kolom baru pada indeks {new_col_idx} di tabel {table_index} pada {path}."
    except Exception as e:
        return f"Error: Gagal menambahkan kolom di {file_path}: {e}"


@mcp.tool()
def resize_table_docx(
    file_path: str,
    table_index: int,
    col_widths: list[float] | None = None,
    col_index: int | None = None,
    width_inches: float | None = None,
) -> str:
    """Mengatur lebar kolom tabel (dalam inci). Bisa mengatur seluruh kolom via col_widths atau satu kolom via col_index & width_inches"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]

        if col_widths is not None:
            for c_idx, w in enumerate(col_widths):
                if c_idx < len(table.columns):
                    w_val = Inches(w)
                    table.columns[c_idx].width = w_val
                    for row in table.rows:
                        if c_idx < len(row.cells):
                            row.cells[c_idx].width = w_val
            doc.save(str(path))
            return f"Berhasil mengatur lebar {min(len(col_widths), len(table.columns))} kolom pada tabel {table_index} di {path}."
        elif col_index is not None and width_inches is not None:
            if col_index < 0 or col_index >= len(table.columns):
                return f"Error: Indeks kolom {col_index} tidak valid. Total kolom: {len(table.columns)}."
            w_val = Inches(width_inches)
            table.columns[col_index].width = w_val
            for row in table.rows:
                if col_index < len(row.cells):
                    row.cells[col_index].width = w_val
            doc.save(str(path))
            return f"Berhasil mengubah lebar kolom {col_index} menjadi {width_inches} inci pada tabel {table_index} di {path}."
        else:
            return "Error: Berikan parameter 'col_widths' (list lebar inci tiap kolom) atau pasangan 'col_index' dan 'width_inches'."
    except Exception as e:
        return f"Error: Gagal mengubah ukuran tabel di {file_path}: {e}"


@mcp.tool()
def set_table_visibility_docx(
    file_path: str,
    table_index: int,
    visible: bool = True,
) -> str:
    """Mengatur visibilitas border tabel (visible=False untuk transparan/tanpa garis, visible=True untuk border garis)"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]
        tblPr = table._element.xpath("w:tblPr")
        if not tblPr:
            return f"Error: Elemen properti tabel {table_index} tidak ditemukan."

        for existing in tblPr[0].xpath("w:tblBorders"):
            tblPr[0].remove(existing)

        val = "single" if visible else "none"
        sz = "4" if visible else "0"
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="auto"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="auto"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="auto"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="auto"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="auto"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="auto"/>'
            f"</w:tblBorders>"
        )
        tblPr[0].append(borders)

        doc.save(str(path))
        status_text = "tampak (bergaris)" if visible else "transparan (tanpa border)"
        return f"Berhasil mengubah visibilitas tabel {table_index} menjadi {status_text} di {path}."
    except Exception as e:
        return f"Error: Gagal mengatur visibilitas tabel di {file_path}: {e}"


@mcp.tool()
def delete_column_docx(
    file_path: str,
    table_index: int,
    col_index: int,
) -> str:
    """Menghapus satu kolom tertentu dari tabel docx berdasarkan col_index (0-indexed)"""
    # shortcut: assume regular unmerged grid; complex merged cells may shift column indices.
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]

        if len(table.columns) <= 1:
            return f"Error: Tabel hanya memiliki 1 kolom. Gunakan delete_table_docx jika ingin menghapus seluruh tabel."

        if col_index < 0 or col_index >= len(table.columns):
            return f"Error: Indeks kolom {col_index} tidak valid. Total kolom: {len(table.columns)}."

        for row in table.rows:
            if col_index < len(row.cells):
                tc = row.cells[col_index]._element
                tc.getparent().remove(tc)

        tblGrid = table._element.xpath("w:tblGrid")
        if tblGrid:
            gridCols = tblGrid[0].xpath("w:gridCol")
            if col_index < len(gridCols):
                tblGrid[0].remove(gridCols[col_index])

        doc.save(str(path))
        return f"Berhasil menghapus kolom {col_index} dari tabel {table_index} di {path}."
    except Exception as e:
        return f"Error: Gagal menghapus kolom di {file_path}: {e}"


@mcp.tool()
def add_row_docx(
    file_path: str,
    table_index: int,
    values: list[str] | None = None,
    height_inches: float | None = None,
) -> str:
    """Menambahkan satu baris baru pada tabel docx dengan teks sel opsional dan tinggi opsional"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]
        row = table.add_row()

        if height_inches and height_inches > 0:
            row.height = Inches(height_inches)

        if values:
            for idx, val in enumerate(values[: len(row.cells)]):
                row.cells[idx].text = str(val)

        doc.save(str(path))
        new_row_idx = len(table.rows) - 1
        return f"Berhasil menambahkan baris baru pada indeks {new_row_idx} di tabel {table_index} pada {path}."
    except Exception as e:
        return f"Error: Gagal menambahkan baris di {file_path}: {e}"


@mcp.tool()
def delete_row_docx(
    file_path: str,
    table_index: int,
    row_index: int,
) -> str:
    """Menghapus satu baris tertentu dari tabel docx berdasarkan row_index (0-indexed)"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]

        if len(table.rows) <= 1:
            return "Error: Tabel hanya memiliki 1 baris. Gunakan delete_table_docx jika ingin menghapus seluruh tabel."

        if row_index < 0 or row_index >= len(table.rows):
            return f"Error: Indeks baris {row_index} tidak valid. Total baris: {len(table.rows)}."

        tr = table.rows[row_index]._element
        tr.getparent().remove(tr)

        doc.save(str(path))
        return f"Berhasil menghapus baris {row_index} dari tabel {table_index} di {path}."
    except Exception as e:
        return f"Error: Gagal menghapus baris di {file_path}: {e}"


@mcp.tool()
def merge_cells_docx(
    file_path: str,
    table_index: int,
    start_row: int,
    start_col: int,
    end_row: int,
    end_col: int,
) -> str:
    """Menggabungkan (merge) area sel dari [start_row, start_col] sampai [end_row, end_col] (0-indexed)"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]

        if not (0 <= start_row < len(table.rows) and 0 <= end_row < len(table.rows)):
            return f"Error: Indeks baris tidak valid (rentang baris tabel: 0 sampai {len(table.rows)-1})."

        if not (0 <= start_col < len(table.columns) and 0 <= end_col < len(table.columns)):
            return f"Error: Indeks kolom tidak valid (rentang kolom tabel: 0 sampai {len(table.columns)-1})."

        cell_a = table.cell(start_row, start_col)
        cell_b = table.cell(end_row, end_col)
        cell_a.merge(cell_b)

        doc.save(str(path))
        return f"Berhasil menggabungkan sel dari [{start_row}, {start_col}] ke [{end_row}, {end_col}] pada tabel {table_index} di {path}."
    except Exception as e:
        return f"Error: Gagal menggabungkan sel di {file_path}: {e}"


@mcp.tool()
def set_cell_shading_docx(
    file_path: str,
    table_index: int,
    color_hex: str,
    row: int | None = None,
    col: int | None = None,
) -> str:
    """Memberikan warna latar belakang (shading) pada sel/baris/kolom/tabel dengan format HEX (contoh: 'E0E0E0' atau '#E0E0E0')"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]
        clean_hex = color_hex.lstrip("#").upper()
        if len(clean_hex) != 6 or not all(c in "0123456789ABCDEF" for c in clean_hex):
            return f"Error: Format warna '{color_hex}' tidak valid. Gunakan format HEX 6 karakter, misal 'E0E0E0' atau '#FFFFFF'."

        def apply_shd(cell):
            tcPr = cell._element.get_or_add_tcPr()
            for s in tcPr.xpath("w:shd"):
                tcPr.remove(s)
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{clean_hex}"/>'))

        if row is not None and col is not None:
            if row < 0 or row >= len(table.rows):
                return f"Error: Indeks baris {row} tidak valid. Total baris: {len(table.rows)}."
            if col < 0 or col >= len(table.rows[row].cells):
                return f"Error: Indeks kolom {col} tidak valid. Total kolom pada baris {row}: {len(table.rows[row].cells)}."
            apply_shd(table.rows[row].cells[col])
            target_desc = f"sel [{row}, {col}]"
        elif row is not None:
            if row < 0 or row >= len(table.rows):
                return f"Error: Indeks baris {row} tidak valid. Total baris: {len(table.rows)}."
            for cell in table.rows[row].cells:
                apply_shd(cell)
            target_desc = f"seluruh baris {row}"
        elif col is not None:
            if col < 0 or col >= len(table.columns):
                return f"Error: Indeks kolom {col} tidak valid. Total kolom: {len(table.columns)}."
            for r in table.rows:
                if col < len(r.cells):
                    apply_shd(r.cells[col])
            target_desc = f"seluruh kolom {col}"
        else:
            for r in table.rows:
                for cell in r.cells:
                    apply_shd(cell)
            target_desc = "seluruh tabel"

        doc.save(str(path))
        return f"Berhasil menerapkan warna #{clean_hex} pada {target_desc} di tabel {table_index} pada {path}."
    except Exception as e:
        return f"Error: Gagal memberikan warna sel di {file_path}: {e}"


@mcp.tool()
def set_table_style_docx(
    file_path: str,
    table_index: int,
    style_name: str,
) -> str:
    """Mengubah template style bawaan Word untuk tabel (misalnya 'Table Grid', 'Light Shading', 'Normal Table')"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]
        table.style = style_name
        doc.save(str(path))
        return f"Berhasil mengubah style tabel {table_index} menjadi '{style_name}' di {path}."
    except Exception as e:
        return f"Error: Gagal mengubah style tabel di {file_path}: {e}"


@mcp.tool()
def set_table_alignment_docx(
    file_path: str,
    table_index: int,
    alignment: str,
) -> str:
    """Mengatur perataan posisi tabel di halaman dokumen ('left', 'center', atau 'right')"""
    try:
        align_map = {
            "left": WD_TABLE_ALIGNMENT.LEFT,
            "center": WD_TABLE_ALIGNMENT.CENTER,
            "right": WD_TABLE_ALIGNMENT.RIGHT,
        }
        align_key = alignment.strip().lower()
        if align_key not in align_map:
            return f"Error: Alignment '{alignment}' tidak valid. Pilihan: 'left', 'center', 'right'."

        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]
        table.alignment = align_map[align_key]
        doc.save(str(path))
        return f"Berhasil mengubah perataan tabel {table_index} menjadi '{align_key}' di {path}."
    except Exception as e:
        return f"Error: Gagal mengatur perataan tabel di {file_path}: {e}"


@mcp.tool()
def set_cell_alignment_docx(
    file_path: str,
    table_index: int,
    horizontal: str = "",
    vertical: str = "",
    row: int | None = None,
    col: int | None = None,
) -> str:
    """Mengatur perataan teks sel: horizontal ('left', 'center', 'right', 'justify') dan/atau vertikal ('top', 'center', 'bottom')"""
    try:
        h_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        v_map = {
            "top": WD_ALIGN_VERTICAL.TOP,
            "center": WD_ALIGN_VERTICAL.CENTER,
            "bottom": WD_ALIGN_VERTICAL.BOTTOM,
        }

        h_val = h_map.get(horizontal.strip().lower()) if horizontal else None
        v_val = v_map.get(vertical.strip().lower()) if vertical else None

        if horizontal and h_val is None:
            return f"Error: Perataan horizontal '{horizontal}' tidak valid. Pilihan: 'left', 'center', 'right', 'justify'."
        if vertical and v_val is None:
            return f"Error: Perataan vertikal '{vertical}' tidak valid. Pilihan: 'top', 'center', 'bottom'."
        if not horizontal and not vertical:
            return "Error: Berikan setidaknya salah satu parameter 'horizontal' atau 'vertical'."

        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]

        def apply_align(cell):
            if v_val is not None:
                cell.vertical_alignment = v_val
            if h_val is not None:
                for p in cell.paragraphs:
                    p.alignment = h_val

        if row is not None and col is not None:
            if row < 0 or row >= len(table.rows):
                return f"Error: Indeks baris {row} tidak valid. Total baris: {len(table.rows)}."
            if col < 0 or col >= len(table.rows[row].cells):
                return f"Error: Indeks kolom {col} tidak valid. Total kolom pada baris {row}: {len(table.rows[row].cells)}."
            apply_align(table.rows[row].cells[col])
            target_desc = f"sel [{row}, {col}]"
        elif row is not None:
            if row < 0 or row >= len(table.rows):
                return f"Error: Indeks baris {row} tidak valid. Total baris: {len(table.rows)}."
            for cell in table.rows[row].cells:
                apply_align(cell)
            target_desc = f"seluruh baris {row}"
        elif col is not None:
            if col < 0 or col >= len(table.columns):
                return f"Error: Indeks kolom {col} tidak valid. Total kolom: {len(table.columns)}."
            for r in table.rows:
                if col < len(r.cells):
                    apply_align(r.cells[col])
            target_desc = f"seluruh kolom {col}"
        else:
            for r in table.rows:
                for cell in r.cells:
                    apply_align(cell)
            target_desc = "seluruh tabel"

        doc.save(str(path))
        return f"Berhasil mengubah perataan pada {target_desc} di tabel {table_index} pada {path}."
    except Exception as e:
        return f"Error: Gagal mengatur perataan sel di {file_path}: {e}"


@mcp.tool()
def set_table_pagination_docx(
    file_path: str,
    table_index: int,
    repeat_header_rows: int = 1,
    cant_split: bool = True,
) -> str:
    """Mengatur opsi cetak/pagination tabel: mengulang header di setiap halaman (repeat_header_rows) dan mencegah baris terbelah (cant_split)"""
    try:
        path = validate_docx_path(file_path, must_exist=True)
        doc = Document(str(path))

        if table_index < 0 or table_index >= len(doc.tables):
            return f"Error: Indeks tabel {table_index} tidak ditemukan. Total tabel: {len(doc.tables)}."

        table = doc.tables[table_index]

        if repeat_header_rows > 0:
            for idx in range(min(repeat_header_rows, len(table.rows))):
                trPr = table.rows[idx]._tr.get_or_add_trPr()
                for h in trPr.xpath("w:tblHeader"):
                    trPr.remove(h)
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        if cant_split:
            for row in table.rows:
                trPr = row._tr.get_or_add_trPr()
                for cs in trPr.xpath("w:cantSplit"):
                    trPr.remove(cs)
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        doc.save(str(path))
        return f"Berhasil mengatur opsi pagination (repeat_header_rows={repeat_header_rows}, cant_split={cant_split}) pada tabel {table_index} di {path}."
    except Exception as e:
        return f"Error: Gagal mengatur pagination tabel di {file_path}: {e}"


if __name__ == "__main__":
    mcp.run()
