import base64
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE, WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START, WD_ORIENTATION
from docx.oxml.ns import qn
from docx_mcp.server import (
    add_column_docx,
    add_equation_docx,
    add_horizontal_rule_docx,
    add_page_break_docx,
    add_picture_docx,
    add_row_docx,
    add_section_docx,
    add_styled_paragraph_docx,
    clear_docx,
    create_table_docx,
    delete_column_docx,
    delete_picture_docx,
    delete_row_docx,
    delete_table_docx,
    edit_table_docx,
    extract_images_docx,
    format_paragraph_docx,
    format_text_docx,
    get_header_footer_docx,
    get_images_info_docx,
    get_page_setup_docx,
    merge_cells_docx,
    read_docx,
    resize_picture_docx,
    resize_table_docx,
    set_cell_alignment_docx,
    set_cell_margins_docx,
    set_cell_shading_docx,
    set_cell_text_direction_docx,
    set_header_footer_docx,
    set_page_setup_docx,
    set_table_alignment_docx,
    set_table_borders_docx,
    set_table_pagination_docx,
    set_table_style_docx,
    set_table_visibility_docx,
    split_table_docx,
    add_table_of_contents_docx,
)


class TestTableTools(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.file_path = str(Path(self.tmp_dir.name) / "test.docx")

    def tearDown(self):
        self.tmp_dir.cleanup()

    def test_create_and_read_table(self):
        data = [
            ["Nama", "Role"],
            ["Alice", "Admin"],
            ["Bob", "User"],
        ]
        res = create_table_docx(self.file_path, data, heading="Daftar Pengguna")
        self.assertIn("Berhasil membuat tabel", res)

        content = read_docx(self.file_path)
        self.assertIn("Daftar Pengguna", content)
        self.assertIn("--- Tabel 0 ---", content)
        self.assertIn("Nama | Role", content)
        self.assertIn("Alice | Admin", content)

    def test_create_table_empty_data(self):
        res = create_table_docx(self.file_path, [])
        self.assertIn("Error: Data tabel tidak boleh kosong", res)

    def test_edit_table(self):
        data = [["Col1", "Col2"], ["Val1", "Val2"]]
        create_table_docx(self.file_path, data)

        res = edit_table_docx(self.file_path, table_index=0, row=1, col=1, new_text="UpdatedVal")
        self.assertIn("Berhasil mengubah sel [1, 1]", res)

        content = read_docx(self.file_path)
        self.assertIn("Val1 | UpdatedVal", content)

    def test_edit_table_invalid_bounds(self):
        data = [["Col1", "Col2"], ["Val1", "Val2"]]
        create_table_docx(self.file_path, data)

        self.assertIn("Error: Indeks tabel 99 tidak ditemukan", edit_table_docx(self.file_path, 99, 0, 0, "X"))
        self.assertIn("Error: Indeks baris 5 tidak valid", edit_table_docx(self.file_path, 0, 5, 0, "X"))
        self.assertIn("Error: Indeks kolom 5 tidak valid", edit_table_docx(self.file_path, 0, 0, 5, "X"))

    def test_delete_table(self):
        data1 = [["T1"]]
        data2 = [["T2"]]
        create_table_docx(self.file_path, data1)
        create_table_docx(self.file_path, data2)

        res = delete_table_docx(self.file_path, table_index=0)
        self.assertIn("Berhasil menghapus tabel 0", res)

        content = read_docx(self.file_path)
        self.assertNotIn("T1", content)
        self.assertIn("T2", content)

    def test_delete_table_invalid_index(self):
        create_table_docx(self.file_path, [["A"]])
        res = delete_table_docx(self.file_path, table_index=3)
        self.assertIn("Error: Indeks tabel 3 tidak ditemukan", res)

    def test_clear_docx(self):
        create_table_docx(self.file_path, [["A", "B"]], heading="Header")
        res = clear_docx(self.file_path)
        self.assertIn("Berhasil mengosongkan seluruh isi file", res)
        self.assertEqual(read_docx(self.file_path), "")

    def test_add_column(self):
        create_table_docx(self.file_path, [["A", "B"], ["1", "2"]])
        res = add_column_docx(self.file_path, table_index=0, values=["C", "3"], width_inches=1.5)
        self.assertIn("Berhasil menambahkan kolom baru", res)
        content = read_docx(self.file_path)
        self.assertIn("A | B | C", content)
        self.assertIn("1 | 2 | 3", content)

    def test_resize_table(self):
        create_table_docx(self.file_path, [["A", "B"]])
        res1 = resize_table_docx(self.file_path, table_index=0, col_widths=[1.5, 2.5])
        self.assertIn("Berhasil mengatur lebar 2 kolom", res1)

        res2 = resize_table_docx(self.file_path, table_index=0, col_index=0, width_inches=3.0)
        self.assertIn("Berhasil mengubah lebar kolom 0 menjadi 3.0 inci", res2)

    def test_set_table_visibility(self):
        create_table_docx(self.file_path, [["A", "B"]])
        res_hide = set_table_visibility_docx(self.file_path, table_index=0, visible=False)
        self.assertIn("transparan (tanpa border)", res_hide)

        res_show = set_table_visibility_docx(self.file_path, table_index=0, visible=True)
        self.assertIn("tampak (bergaris)", res_show)

    def test_delete_column(self):
        create_table_docx(self.file_path, [["A", "B", "C"], ["1", "2", "3"]])
        res = delete_column_docx(self.file_path, table_index=0, col_index=1)
        self.assertIn("Berhasil menghapus kolom 1", res)

        content = read_docx(self.file_path)
        self.assertIn("A | C", content)
        self.assertIn("1 | 3", content)

    def test_delete_column_last_remaining(self):
        create_table_docx(self.file_path, [["A"]])
        res = delete_column_docx(self.file_path, table_index=0, col_index=0)
        self.assertIn("Error: Tabel hanya memiliki 1 kolom", res)

    def test_add_row(self):
        create_table_docx(self.file_path, [["A", "B"]])
        res = add_row_docx(self.file_path, table_index=0, values=["C", "D"], height_inches=0.5)
        self.assertIn("Berhasil menambahkan baris baru pada indeks 1", res)
        content = read_docx(self.file_path)
        self.assertIn("C | D", content)

    def test_delete_row(self):
        create_table_docx(self.file_path, [["R0C0"], ["R1C0"], ["R2C0"]])
        res = delete_row_docx(self.file_path, table_index=0, row_index=1)
        self.assertIn("Berhasil menghapus baris 1", res)
        content = read_docx(self.file_path)
        self.assertIn("R0C0", content)
        self.assertNotIn("R1C0", content)
        self.assertIn("R2C0", content)

    def test_delete_row_last_remaining(self):
        create_table_docx(self.file_path, [["A"]])
        res = delete_row_docx(self.file_path, table_index=0, row_index=0)
        self.assertIn("Error: Tabel hanya memiliki 1 baris", res)

    def test_merge_cells(self):
        create_table_docx(self.file_path, [["A", "B"], ["C", "D"]])
        res = merge_cells_docx(self.file_path, table_index=0, start_row=0, start_col=0, end_row=0, end_col=1)
        self.assertIn("Berhasil menggabungkan sel dari [0, 0] ke [0, 1]", res)

    def test_set_cell_shading(self):
        create_table_docx(self.file_path, [["A", "B"], ["C", "D"]])
        res_cell = set_cell_shading_docx(self.file_path, table_index=0, color_hex="#CCCCCC", row=0, col=0)
        self.assertIn("Berhasil menerapkan warna #CCCCCC pada sel [0, 0]", res_cell)

        res_row = set_cell_shading_docx(self.file_path, table_index=0, color_hex="E0E0E0", row=1)
        self.assertIn("Berhasil menerapkan warna #E0E0E0 pada seluruh baris 1", res_row)

        res_invalid = set_cell_shading_docx(self.file_path, table_index=0, color_hex="INVALID")
        self.assertIn("Error: Format warna 'INVALID' tidak valid", res_invalid)

    def test_set_table_style(self):
        create_table_docx(self.file_path, [["A", "B"]])
        res = set_table_style_docx(self.file_path, table_index=0, style_name="Table Grid")
        self.assertIn("Berhasil mengubah style tabel 0", res)

    def test_set_table_alignment(self):
        create_table_docx(self.file_path, [["A", "B"]])
        res = set_table_alignment_docx(self.file_path, table_index=0, alignment="center")
        self.assertIn("Berhasil mengubah perataan tabel 0 menjadi 'center'", res)

    def test_set_cell_alignment(self):
        create_table_docx(self.file_path, [["A", "B"], ["C", "D"]])
        res = set_cell_alignment_docx(self.file_path, table_index=0, horizontal="center", vertical="center", row=0, col=0)
        self.assertIn("Berhasil mengubah perataan pada sel [0, 0]", res)

    def test_set_table_pagination(self):
        create_table_docx(self.file_path, [["Header 1", "Header 2"], ["Data 1", "Data 2"]])
        res = set_table_pagination_docx(self.file_path, table_index=0, repeat_header_rows=1, cant_split=True)
        self.assertIn("Berhasil mengatur opsi pagination", res)

    def test_insert_row_and_column_at_index(self):
        create_table_docx(self.file_path, [["A", "C"], ["1", "3"]])
        res_col = add_column_docx(self.file_path, table_index=0, values=["B", "2"], insert_index=1)
        self.assertIn("Berhasil menambahkan kolom baru pada indeks 1", res_col)
        content1 = read_docx(self.file_path)
        self.assertIn("A | B | C", content1)

        res_row = add_row_docx(self.file_path, table_index=0, values=["X", "Y", "Z"], insert_index=1)
        self.assertIn("Berhasil menambahkan baris baru pada indeks 1", res_row)
        content2 = read_docx(self.file_path)
        self.assertIn("X | Y | Z", content2)

    def test_set_cell_margins(self):
        create_table_docx(self.file_path, [["A", "B"]])
        res_table = set_cell_margins_docx(self.file_path, table_index=0, top_pt=5.0, bottom_pt=5.0)
        self.assertIn("Berhasil mengatur margin padding", res_table)

        res_cell = set_cell_margins_docx(self.file_path, table_index=0, row=0, col=0, left_pt=10.0)
        self.assertIn("pada sel [0, 0]", res_cell)

    def test_set_table_borders(self):
        create_table_docx(self.file_path, [["A", "B"]])
        res = set_table_borders_docx(self.file_path, table_index=0, color_hex="FF0000", size=8, top=False, bottom=True)
        self.assertIn("Berhasil mengatur border kustom", res)

    def test_set_cell_text_direction(self):
        create_table_docx(self.file_path, [["A", "B"]])
        res = set_cell_text_direction_docx(self.file_path, table_index=0, direction="bt_lr", row=0, col=0)
        self.assertIn("Berhasil mengubah arah teks menjadi 'bt_lr' pada sel [0, 0]", res)

    def test_split_table(self):
        create_table_docx(self.file_path, [["R0"], ["R1"], ["R2"], ["R3"]])
        res = split_table_docx(self.file_path, table_index=0, split_row_index=2)
        self.assertIn("Berhasil membagi tabel 0 pada baris 2", res)
        self.assertIn("memiliki 2 tabel", res)

    def test_add_styled_paragraph(self):
        res = add_styled_paragraph_docx(
            self.file_path,
            text="Heading Bergaya",
            heading_level=1,
            font_name="Times New Roman",
            font_size_pt=18.0,
            color="#003366",
            bold=True,
            italic=True,
            alignment="center",
        )
        self.assertIn("Berhasil menambahkan paragraf berformat", res)

    def test_format_text_paragraph(self):
        add_styled_paragraph_docx(self.file_path, text="Teks awal sebelum diformat.")
        res = format_text_docx(
            self.file_path,
            paragraph_index=0,
            font_name="Arial",
            font_size_pt=12.0,
            color="red",
            bold=True,
            underline=True,
            strike=True,
            highlight_color="yellow",
        )
        self.assertIn("Berhasil memformat teks pada paragraf 0", res)
        self.assertIn("font='Arial'", res)
        self.assertIn("size=12.0pt", res)

    def test_format_text_target_phrase(self):
        add_styled_paragraph_docx(self.file_path, text="Kalimat ini ada kata kunci penting di dalamnya.")
        res = format_text_docx(
            self.file_path,
            target_text="kata kunci",
            font_name="Courier New",
            font_size_pt=14.0,
            color="#FF5500",
            bold=True,
        )
        self.assertIn("Berhasil memformat teks pada seluruh dokumen untuk teks 'kata kunci'", res)

    def test_format_text_table_cell(self):
        create_table_docx(self.file_path, [["Nama", "Skor"], ["Budi", "100"]])
        res = format_text_docx(
            self.file_path,
            table_index=0,
            row=0,
            font_name="Arial",
            bold=True,
            color="#FFFFFF",
        )
        self.assertIn("tabel 0 baris 0", res)

    def test_format_text_typography(self):
        add_styled_paragraph_docx(self.file_path, text="H2O formula and special text")
        res = format_text_docx(
            self.file_path,
            paragraph_index=0,
            subscript=True,
            small_caps=True,
            underline_style="double",
        )
        self.assertIn("subscript=True", res)
        self.assertIn("small_caps=True", res)
        self.assertIn("underline_style='double'", res)

        doc = Document(self.file_path)
        run = doc.paragraphs[0].runs[0]
        self.assertTrue(run.font.subscript)
        self.assertTrue(run.font.small_caps)
        self.assertEqual(run.font.underline, WD_UNDERLINE.DOUBLE)

    def test_format_paragraph(self):
        add_styled_paragraph_docx(self.file_path, text="Paragraf uji pemformatan spasi.")
        res = format_paragraph_docx(
            self.file_path,
            paragraph_index=0,
            line_spacing=1.5,
            space_before_pt=6.0,
            space_after_pt=12.0,
            first_line_indent_inches=0.5,
            left_indent_inches=0.25,
            alignment="justify",
            list_style="bullet",
        )
        self.assertIn("Berhasil memformat tata letak dan spasi pada paragraf 0", res)

        doc = Document(self.file_path)
        p = doc.paragraphs[0]
        self.assertEqual(p.paragraph_format.line_spacing, 1.5)
        self.assertAlmostEqual(p.paragraph_format.space_before.pt, 6.0)
        self.assertAlmostEqual(p.paragraph_format.space_after.pt, 12.0)
        self.assertAlmostEqual(p.paragraph_format.first_line_indent.inches, 0.5)
        self.assertAlmostEqual(p.paragraph_format.left_indent.inches, 0.25)
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

    def test_add_styled_paragraph_advanced(self):
        res = add_styled_paragraph_docx(
            self.file_path,
            text="E=mc2 formula",
            line_spacing=2.0,
            space_before_pt=4.0,
            space_after_pt=10.0,
            first_line_indent_inches=0.5,
            left_indent_inches=0.3,
            right_indent_inches=0.2,
            superscript=True,
            all_caps=True,
            underline_style="dotted",
            alignment="center",
        )
        self.assertIn("Berhasil menambahkan paragraf berformat", res)

        doc = Document(self.file_path)
        p = doc.paragraphs[0]
        self.assertEqual(p.paragraph_format.line_spacing, 2.0)
        self.assertAlmostEqual(p.paragraph_format.space_before.pt, 4.0)
        self.assertAlmostEqual(p.paragraph_format.space_after.pt, 10.0)
        self.assertAlmostEqual(p.paragraph_format.first_line_indent.inches, 0.5)
        self.assertAlmostEqual(p.paragraph_format.left_indent.inches, 0.3)
        self.assertAlmostEqual(p.paragraph_format.right_indent.inches, 0.2)
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)

        run = p.runs[0]
        self.assertTrue(run.font.superscript)
        self.assertTrue(run.font.all_caps)
        self.assertEqual(run.font.underline, WD_UNDERLINE.DOTTED)

    def test_format_paragraph_table_and_flow_control(self):
        create_table_docx(self.file_path, [["H1", "H2"], ["A", "B"]])
        res = format_paragraph_docx(
            self.file_path,
            table_index=0,
            row=1,
            col=1,
            line_spacing_pt=14.0,
            alignment="center",
        )
        self.assertIn("tabel 0 sel [1, 1]", res)

        doc = Document(self.file_path)
        p = doc.tables[0].rows[1].cells[1].paragraphs[0]
        self.assertAlmostEqual(p.paragraph_format.line_spacing.pt, 14.0)
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_format_paragraph_hanging_and_flow(self):
        add_styled_paragraph_docx(self.file_path, text="Daftar pustaka referensi...")
        res = format_paragraph_docx(
            self.file_path,
            paragraph_index=0,
            hanging_indent_inches=0.5,
            keep_with_next=True,
            page_break_before=True,
            widow_control=True,
        )
        self.assertIn("Berhasil memformat tata letak dan spasi pada paragraf 0", res)

        doc = Document(self.file_path)
        p = doc.paragraphs[0]
        self.assertAlmostEqual(p.paragraph_format.left_indent.inches, 0.5)
        self.assertAlmostEqual(p.paragraph_format.first_line_indent.inches, -0.5)
        self.assertTrue(p.paragraph_format.keep_with_next)
        self.assertTrue(p.paragraph_format.page_break_before)
        self.assertTrue(p.paragraph_format.widow_control)

    def test_add_picture_basic_and_read(self):
        img_path = Path(self.tmp_dir.name) / "test.png"
        img_path.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="))

        res = add_picture_docx(
            self.file_path,
            str(img_path),
            width_inches=2.0,
            height_inches=1.5,
            alignment="center",
            caption="Gambar 1: Sampel Visual",
        )
        self.assertIn("Berhasil menyisipkan gambar 'test.png'", res)

        content = read_docx(self.file_path)
        self.assertIn("Total Gambar: 1", content)
        self.assertIn("Gambar 1: Sampel Visual", content)

        doc = Document(self.file_path)
        self.assertEqual(len(doc.inline_shapes), 1)
        shape = doc.inline_shapes[0]
        self.assertAlmostEqual(shape.width / 914400.0, 2.0, places=1)
        self.assertAlmostEqual(shape.height / 914400.0, 1.5, places=1)

    def test_add_picture_table_cell_and_cm(self):
        img_path = Path(self.tmp_dir.name) / "test_table.png"
        img_path.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="))

        create_table_docx(self.file_path, [["Produk", "Foto"]])
        res = add_picture_docx(
            self.file_path,
            str(img_path),
            table_index=0,
            row=0,
            col=1,
            width_cm=4.0,
            alignment="center",
        )
        self.assertIn("tabel 0 sel [0, 1]", res)

        doc = Document(self.file_path)
        cell = doc.tables[0].rows[0].cells[1]
        self.assertEqual(cell.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_get_images_info_and_resize(self):
        img_path = Path(self.tmp_dir.name) / "test_resize.png"
        img_path.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="))

        add_picture_docx(self.file_path, str(img_path), width_inches=1.0)
        info = get_images_info_docx(self.file_path)
        self.assertIn("Total gambar pada", info)
        self.assertIn("[Gambar 0]", info)

        resize_res = resize_picture_docx(self.file_path, picture_index=0, width_inches=3.0, height_inches=2.0)
        self.assertIn("Berhasil mengubah ukuran gambar 0", resize_res)

        doc = Document(self.file_path)
        shape = doc.inline_shapes[0]
        self.assertAlmostEqual(shape.width / 914400.0, 3.0, places=1)
        self.assertAlmostEqual(shape.height / 914400.0, 2.0, places=1)

    def test_add_picture_base64_and_alt_text(self):
        b64_data = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        res = add_picture_docx(
            self.file_path,
            image_base64=b64_data,
            width_inches=1.5,
            alt_text="Logo Perusahaan",
            title="Logo Utama",
        )
        self.assertIn("Berhasil menyisipkan gambar base64", res)

        doc = Document(self.file_path)
        self.assertEqual(len(doc.inline_shapes), 1)
        shape = doc.inline_shapes[0]
        self.assertEqual(shape._inline.docPr.get("descr"), "Logo Perusahaan")
        self.assertEqual(shape._inline.docPr.get("title"), "Logo Utama")

    def test_delete_picture(self):
        b64_data = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        add_picture_docx(self.file_path, image_base64=b64_data)
        doc = Document(self.file_path)
        self.assertEqual(len(doc.inline_shapes), 1)

        res = delete_picture_docx(self.file_path, picture_index=0)
        self.assertIn("Berhasil menghapus gambar 0", res)

        doc_after = Document(self.file_path)
        self.assertEqual(len(doc_after.inline_shapes), 0)

    def test_extract_images(self):
        b64_data = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        add_picture_docx(self.file_path, image_base64=b64_data)

        out_dir = Path(self.tmp_dir.name) / "extracted_imgs"
        res = extract_images_docx(self.file_path, str(out_dir))
        self.assertIn("Berhasil mengekstrak 1 gambar", res)
        self.assertTrue(out_dir.is_dir())
        files = list(out_dir.glob("*"))
        self.assertEqual(len(files), 1)

    def test_set_page_setup_and_read(self):
        res = set_page_setup_docx(
            self.file_path,
            paper_size="a4",
            orientation="portrait",
            margin_preset="skripsi",
            gutter_cm=0.5,
        )
        self.assertIn("Berhasil memperbarui pengaturan halaman", res)

        info = get_page_setup_docx(self.file_path)
        self.assertIn("A4", info)
        self.assertIn("Portrait", info)
        self.assertIn("Atas=4.00cm", info)
        self.assertIn("Gutter: 0.50 cm", info)

        doc = Document(self.file_path)
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.top_margin / 360000.0, 4.0, places=1)
        self.assertAlmostEqual(sec.left_margin / 360000.0, 4.0, places=1)
        self.assertAlmostEqual(sec.bottom_margin / 360000.0, 3.0, places=1)
        self.assertAlmostEqual(sec.right_margin / 360000.0, 3.0, places=1)
        self.assertAlmostEqual(sec.page_width / 360000.0, 21.0, places=1)
        self.assertAlmostEqual(sec.page_height / 360000.0, 29.7, places=1)

    def test_set_page_setup_landscape_and_custom_margins(self):
        res = set_page_setup_docx(
            self.file_path,
            paper_size="letter",
            orientation="landscape",
            top_margin_inches=1.0,
            bottom_margin_inches=1.0,
            left_margin_inches=1.5,
            right_margin_inches=1.5,
        )
        self.assertIn("Berhasil memperbarui pengaturan halaman", res)

        doc = Document(self.file_path)
        sec = doc.sections[0]
        self.assertEqual(sec.orientation, WD_ORIENTATION.LANDSCAPE)
        self.assertAlmostEqual(sec.page_width / 914400.0, 11.0, places=1)
        self.assertAlmostEqual(sec.page_height / 914400.0, 8.5, places=1)
        self.assertAlmostEqual(sec.left_margin / 914400.0, 1.5, places=1)

    def test_add_section_and_multi_section(self):
        set_page_setup_docx(self.file_path, paper_size="a4", orientation="portrait")
        res = add_section_docx(
            self.file_path,
            start_type="new_page",
            paper_size="a4",
            orientation="landscape",
            margin_preset="narrow",
        )
        self.assertIn("Berhasil menambahkan section baru", res)

        doc = Document(self.file_path)
        self.assertEqual(len(doc.sections), 2)
        self.assertEqual(doc.sections[0].orientation, WD_ORIENTATION.PORTRAIT)
        self.assertEqual(doc.sections[1].orientation, WD_ORIENTATION.LANDSCAPE)
        self.assertAlmostEqual(doc.sections[1].top_margin / 360000.0, 1.27, places=1)

    def test_page_setup_columns_and_numbering(self):
        res = set_page_setup_docx(
            self.file_path,
            columns_count=2,
            column_spacing_cm=1.0,
            page_number_start=1,
            page_number_format="lowerRoman",
            mirror_margins=True,
        )
        self.assertIn("Berhasil memperbarui pengaturan halaman", res)

        info = get_page_setup_docx(self.file_path)
        self.assertIn("Jumlah Kolom: 2", info)
        self.assertIn("Mulai=1, Format=lowerRoman", info)
        self.assertIn("Status Margin Cermin (Mirror Margins): Aktif", info)

        doc = Document(self.file_path)
        sec = doc.sections[0]
        cols = sec._sectPr.find(qn("w:cols"))
        self.assertIsNotNone(cols)
        self.assertEqual(cols.get(qn("w:num")), "2")

        pg = sec._sectPr.find(qn("w:pgNumType"))
        self.assertIsNotNone(pg)
        self.assertEqual(pg.get(qn("w:fmt")), "lowerRoman")
        self.assertEqual(pg.get(qn("w:start")), "1")

        mirror = doc.settings.element.find(qn("w:mirrorMargins"))
        self.assertIsNotNone(mirror)

    def test_add_section_with_columns(self):
        add_section_docx(
            self.file_path,
            start_type="new_page",
            columns_count=3,
            column_spacing_cm=0.5,
        )
        doc = Document(self.file_path)
        self.assertEqual(len(doc.sections), 2)
        cols = doc.sections[1]._sectPr.find(qn("w:cols"))
        self.assertIsNotNone(cols)
        self.assertEqual(cols.get(qn("w:num")), "3")

    def test_add_page_break(self):
        doc = Document()
        doc.add_paragraph("Paragraf 1")
        doc.add_paragraph("Paragraf 2")
        doc.save(self.file_path)

        # 1. Break di akhir
        res1 = add_page_break_docx(self.file_path, break_type="page")
        self.assertIn("Berhasil menyisipkan page break di akhir", res1)

        # 2. Break sebelum paragraf 1
        res2 = add_page_break_docx(self.file_path, break_type="column", paragraph_index=1)
        self.assertIn("Berhasil menyisipkan column break sebelum paragraf 1", res2)

        # 3. Break di dalam sel tabel
        create_table_docx(self.file_path, [["A", "B"], ["C", "D"]])
        res3 = add_page_break_docx(self.file_path, break_type="page", table_index=0, row=0, col=1)
        self.assertIn("Berhasil menyisipkan page break di dalam tabel 0 sel [0, 1]", res3)

        doc_check = Document(self.file_path)
        # Verifikasi bahwa run memiliki break
        breaks = doc_check.element.xpath(".//w:br")
        self.assertGreaterEqual(len(breaks), 3)

        # Verifikasi bahwa read_docx menampilkan marker page break
        content = read_docx(self.file_path)
        self.assertIn("[Page Break]", content)

    def test_add_horizontal_rule(self):
        doc = Document()
        doc.add_paragraph("Kementerian Pendidikan dan Kebudayaan")
        doc.add_paragraph("Universitas Indonesia")
        doc.add_paragraph("Jl. Salemba Raya No. 4, Jakarta Pusat")
        doc.save(self.file_path)

        # 1. Terapkan garis kop surat instansi resmi (thick_thin) langsung pada paragraf alamat
        res_kop = add_horizontal_rule_docx(
            self.file_path,
            style="thick_thin",
            size_pt=2.25,
            color="#000000",
            space_after_pt=4.0,
            paragraph_index=2,
            apply_to_existing=True,
        )
        self.assertIn("Berhasil menambahkan garis pembatas", res_kop)

        # 2. Sisipkan standalone rule pembatas bab dengan indentasi
        res_rule = add_horizontal_rule_docx(
            self.file_path,
            style="single",
            size_pt=1.5,
            color="blue",
            space_before_pt=6.0,
            space_after_pt=12.0,
            indent_left_inches=1.0,
            indent_right_inches=1.0,
            alignment="center",
        )
        self.assertIn("Berhasil menambahkan garis pembatas", res_rule)

        doc_check = Document(self.file_path)
        # Cek border pada paragraf kop surat
        p_kop = doc_check.paragraphs[2]
        pPr_kop = p_kop._p.find(qn("w:pPr"))
        self.assertIsNotNone(pPr_kop)
        pBdr_kop = pPr_kop.find(qn("w:pBdr"))
        self.assertIsNotNone(pBdr_kop)
        bottom_kop = pBdr_kop.find(qn("w:bottom"))
        self.assertIsNotNone(bottom_kop)
        self.assertEqual(bottom_kop.get(qn("w:val")), "thickThinMediumGap")
        self.assertEqual(bottom_kop.get(qn("w:sz")), "18")  # 2.25 * 8 = 18

        # Cek standalone rule di akhir dokumen
        p_last = doc_check.paragraphs[-1]
        self.assertAlmostEqual(p_last.paragraph_format.left_indent.inches, 1.0, places=1)
        self.assertAlmostEqual(p_last.paragraph_format.right_indent.inches, 1.0, places=1)
        pPr_last = p_last._p.find(qn("w:pPr"))
        self.assertIsNotNone(pPr_last)
        pBdr_last = pPr_last.find(qn("w:pBdr"))
        self.assertIsNotNone(pBdr_last)
        bottom_last = pBdr_last.find(qn("w:bottom"))
        self.assertIsNotNone(bottom_last)
        self.assertEqual(bottom_last.get(qn("w:val")), "single")
        self.assertEqual(bottom_last.get(qn("w:color")), "0000FF")  # blue

        # Cek visibilitas di read_docx
        content = read_docx(self.file_path)
        self.assertIn("[Garis Pembatas: thickThinMediumGap]", content)
        self.assertIn("--- Garis Horizontal (single) ---", content)

    def test_add_equation_block_latex(self):
        # 1. Tambah rumus kuadrat mode block
        latex = r"x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}"
        res = add_equation_docx(self.file_path, equation=latex, display_mode="block")
        self.assertIn("Berhasil menyisipkan rumus matematika", res)

        doc = Document(self.file_path)
        p = doc.paragraphs[0]
        # Verifikasi keberadaan tag oMath dan oMathPara
        self.assertIsNotNone(p._p.find(qn("m:oMathPara")))
        self.assertIsNotNone(p._p.find(qn("m:oMathPara")).find(qn("m:oMath")))

        # Verifikasi pembacaan lewat read_docx
        content = read_docx(self.file_path)
        self.assertIn("[Rumus:", content)

    def test_add_equation_inline_and_table(self):
        doc = Document()
        doc.add_paragraph("Teks awal:")
        doc.save(self.file_path)

        # 1. Rumus inline dengan text_before dan text_after
        res_inline = add_equation_docx(
            self.file_path,
            equation=r"E = mc^2",
            display_mode="inline",
            text_before="Teori relativitas: ",
            text_after=" adalah rumus terkenal Einstein.",
            paragraph_index=0,
            insert_before=False,
        )
        self.assertIn("Berhasil menyisipkan rumus matematika", res_inline)

        # 2. Rumus di dalam tabel
        create_table_docx(self.file_path, [["No", "Rumus"]])
        res_tbl = add_equation_docx(
            self.file_path,
            equation=r"\int_0^1 x dx",
            table_index=0,
            row=0,
            col=1,
            display_mode="inline",
        )
        self.assertIn("Berhasil menyisipkan rumus matematika (inline) di dalam tabel 0 sel [0, 1]", res_tbl)

        content = read_docx(self.file_path)
        self.assertIn("Teori relativitas:", content)
        self.assertIn("[Rumus:", content)
        self.assertIn("adalah rumus terkenal Einstein.", content)

    def test_add_equation_mathml_and_error(self):
        # 1. MathML input
        mathml = '<math xmlns="http://www.w3.org/1998/Math/MathML"><mi>a</mi><mo>+</mo><mi>b</mi></math>'
        res_mathml = add_equation_docx(self.file_path, equation=mathml, display_mode="inline")
        self.assertIn("Berhasil menyisipkan rumus matematika", res_mathml)

        # 2. Input kosong
        res_err_empty = add_equation_docx(self.file_path, equation="")
        self.assertIn("Error: Parameter equation tidak boleh kosong", res_err_empty)

        # 3. Input sintaks latex tidak valid
        res_err_syntax = add_equation_docx(self.file_path, equation=r"\frac{a}{")
        self.assertIn("Error: Gagal mengonversi rumus matematika", res_err_syntax)

    def test_add_equation_numbered_and_placeholder(self):
        doc = Document()
        doc.add_paragraph("Solusi dari sistem adalah {{eq_placeholder}} dengan batas x > 0.")
        doc.save(self.file_path)

        # 1. Ganti placeholder dengan equation dan custom font_size_pt
        res_ph = add_equation_docx(
            self.file_path,
            equation=r"x = \sqrt{y}",
            placeholder="{{eq_placeholder}}",
            font_size_pt=14.0,
        )
        self.assertIn("Berhasil mengganti placeholder '{{eq_placeholder}}'", res_ph)

        content = read_docx(self.file_path)
        self.assertNotIn("{{eq_placeholder}}", content)
        self.assertIn("Solusi dari sistem adalah", content)
        self.assertIn("[Rumus:", content)
        self.assertIn("dengan batas x > 0.", content)

        # 2. Rumus bernomor ala jurnal ilmiah
        res_num = add_equation_docx(
            self.file_path,
            equation=r"F = ma",
            equation_number="(2.1)",
            font_size_pt=12.0,
            space_before_pt=6.0,
            space_after_pt=6.0,
        )
        self.assertIn("dengan nomor '(2.1)'", res_num)

        doc_check = Document(self.file_path)
        p_num = doc_check.paragraphs[-1]
        tabs = p_num._p.xpath(".//w:tabs/w:tab")
        self.assertEqual(len(tabs), 2)
        r_texts = [r.text for r in p_num.runs]
        self.assertIn("(2.1)", "".join(r_texts))

    def test_set_header_footer_skripsi_chapter_preset(self):
        doc = Document()
        doc.add_paragraph("Bab I Pendahuluan")
        doc.save(self.file_path)

        res = set_header_footer_docx(
            self.file_path,
            preset="skripsi_chapter",
            text="BAB I PENDAHULUAN",
            font_name="Times New Roman",
            font_size_pt=12.0,
        )
        self.assertIn("Berhasil mengatur Header/Footer (preset 'skripsi_chapter')", res)

        doc_check = Document(self.file_path)
        sec = doc_check.sections[0]
        self.assertTrue(sec.different_first_page_header_footer)

        # First page footer has center page number
        f_p = sec.first_page_footer.paragraphs[0]
        self.assertEqual(f_p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(bool(f_p._p.xpath('.//w:fldSimple[@w:instr="PAGE"]')))

        # Normal header has right page number
        h_p = sec.header.paragraphs[0]
        self.assertEqual(h_p.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertTrue(bool(h_p._p.xpath('.//w:fldSimple[@w:instr="PAGE"]')))
        self.assertIn("BAB I PENDAHULUAN", [r.text for r in h_p.runs])

        # Test inspection tool
        info = get_header_footer_docx(self.file_path)
        self.assertIn("Different First Page: Ya", info)
        self.assertIn("[PAGE]", info)

    def test_set_header_footer_custom_format_and_unlink(self):
        doc = Document()
        doc.add_paragraph("Section 1")
        doc.add_section()
        doc.add_paragraph("Section 2")
        doc.save(self.file_path)

        # Format custom dengan template Halaman {page} dari {pages}
        res = set_header_footer_docx(
            self.file_path,
            target="footer",
            section_index=1,
            is_linked_to_previous=False,
            page_number_format_text="Halaman {page} dari {pages}",
            alignment="center",
            font_name="Arial",
            font_size_pt=10.0,
            color="#0000FF",
        )
        self.assertIn("Berhasil mengatur Header/Footer", res)

        doc_check = Document(self.file_path)
        sec2 = doc_check.sections[1]
        self.assertFalse(sec2.footer.is_linked_to_previous)

        p = sec2.footer.paragraphs[0]
        fld_page = p._p.xpath('.//w:fldSimple[@w:instr="PAGE"]')
        fld_numpages = p._p.xpath('.//w:fldSimple[@w:instr="NUMPAGES"]')
        self.assertTrue(bool(fld_page))
        self.assertTrue(bool(fld_numpages))

        info = get_header_footer_docx(self.file_path, section_index=1)
        self.assertIn("[PAGE]", info)
        self.assertIn("[NUMPAGES]", info)
        self.assertIn("Unlinked", info)

    def test_add_table_of_contents_dynamic(self):
        doc = Document()
        doc.add_paragraph("Introductory text")
        doc.save(self.file_path)

        res = add_table_of_contents_docx(
            self.file_path,
            heading="DAFTAR ISI",
            toc_type="dynamic",
            max_level=3,
            page_break_after=True,
            update_fields_on_open=True,
        )
        self.assertIn("Berhasil menambahkan 'DAFTAR ISI'", res)

        doc_check = Document(self.file_path)
        self.assertEqual(doc_check.paragraphs[1].text, "DAFTAR ISI")
        # Check dynamic TOC field
        toc_fld = doc_check.paragraphs[2]._p.xpath('.//w:fldSimple[contains(@w:instr, "TOC")]')
        self.assertTrue(bool(toc_fld))
        instr = toc_fld[0].get(qn("w:instr"))
        self.assertIn(r'TOC \o "1-3" \h \z \u', instr)
        # Check page break after
        br_p = doc_check.paragraphs[3]
        self.assertTrue(bool(br_p._p.xpath('.//w:br[@w:type="page"]')))
        # Check updateFields setting
        upd = doc_check.settings.element.find(qn("w:updateFields"))
        self.assertIsNotNone(upd)
        self.assertEqual(upd.get(qn("w:val")), "true")

        # Check read_docx recognition
        read_output = read_docx(self.file_path)
        self.assertIn("[TOC Field]", read_output)

    def test_add_table_of_contents_manual_entries(self):
        doc = Document()
        doc.save(self.file_path)

        entries = [
            {"title": "LEMBAR PENGESAHAN", "page": "ii", "level": 1},
            {"title": "KATA PENGANTAR", "page": "iii", "level": 1},
            {"title": "BAB I PENDAHULUAN", "page": "1", "level": 1},
            {"title": "1.1 Latar Belakang", "page": "1", "level": 2},
            {"title": "1.2 Rumusan Masalah", "page": "3", "level": 2},
        ]
        res = add_table_of_contents_docx(
            self.file_path,
            heading="DAFTAR ISI",
            toc_type="static",
            entries=entries,
            tab_leader="dots",
            level_indent_cm=0.5,
            line_spacing=1.5,
            page_break_after=True,
        )
        self.assertIn("5 entri", res)

        doc_check = Document(self.file_path)
        # Heading
        self.assertEqual(doc_check.paragraphs[0].text, "DAFTAR ISI")
        # Entry 1 (level 1 bold)
        p1 = doc_check.paragraphs[1]
        self.assertIn("LEMBAR PENGESAHAN", p1.text)
        self.assertIn("ii", p1.text)
        self.assertTrue(p1.runs[0].bold)
        # Tab stops
        self.assertGreater(len(p1.paragraph_format.tab_stops), 0)
        ts = p1.paragraph_format.tab_stops[0]
        self.assertEqual(ts.leader, WD_TAB_LEADER.DOTS)
        # Entry 4 (level 2 indent)
        p4 = doc_check.paragraphs[4]
        self.assertIn("1.1 Latar Belakang", p4.text)
        self.assertGreater(p4.paragraph_format.left_indent, 0)
        self.assertFalse(p4.runs[0].bold)

    def test_add_table_of_contents_auto_headings_and_position(self):
        doc = Document()
        doc.add_heading("Bab I Pendahuluan", level=1)
        doc.add_paragraph("Isi bab satu...")
        doc.add_heading("1.1 Latar Belakang", level=2)
        doc.add_paragraph("Isi latar belakang...")
        doc.save(self.file_path)

        # Insert TOC at index 0 (before paragraph 0)
        res = add_table_of_contents_docx(
            self.file_path,
            heading="DAFTAR ISI",
            toc_type="auto",
            paragraph_index=0,
            insert_before=True,
            page_break_after=True,
        )
        self.assertIn("Berhasil menambahkan 'DAFTAR ISI'", res)

        doc_check = Document(self.file_path)
        # First paragraph must now be DAFTAR ISI
        self.assertEqual(doc_check.paragraphs[0].text, "DAFTAR ISI")
        # Auto-extracted entries
        self.assertIn("Bab I Pendahuluan", doc_check.paragraphs[1].text)
        self.assertIn("1.1 Latar Belakang", doc_check.paragraphs[2].text)
        # Original chapter is now pushed down after TOC and page break
        self.assertTrue(any("Isi bab satu" in p.text for p in doc_check.paragraphs))

    def test_add_table_of_contents_custom_title_and_error(self):
        doc = Document()
        doc.add_paragraph("Existing")
        doc.save(self.file_path)

        # Custom heading e.g. DAFTAR TABEL, without page break after
        res = add_table_of_contents_docx(
            self.file_path,
            heading="DAFTAR TABEL",
            toc_type="dynamic",
            page_break_after=False,
            page_break_before=True,
        )
        self.assertIn("Berhasil menambahkan 'DAFTAR TABEL'", res)

        # Error cases
        err_idx = add_table_of_contents_docx(self.file_path, paragraph_index=999)
        self.assertIn("Error: Indeks paragraf 999 tidak ditemukan", err_idx)

        err_file = add_table_of_contents_docx("non_existent_file.docx")
        self.assertIn("Error: Gagal menambahkan Table of Contents", err_file)


if __name__ == "__main__":
    unittest.main()






