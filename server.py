from mcp.server.mcpserver import MCPServer
from docx import Document
import os

mcp = MCPServer("DocxEditor")

@mcp.tool()
def read_docx(file_path: str) -> str:
    """Membaca seluruh isi teks dan tabel dari file .docx"""
    if not os.path.exists(file_path):
        return f"Error: File {file_path} tidak ditemukan."

    try:
        doc = Document(file_path)
    except Exception as e:
        return f"Error: Gagal membaca file {file_path}: {e}"

    content = []

    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            content.append(f"[P{i}] {p.text}")

    for t_idx, table in enumerate(doc.tables):
        content.append(f"\n--- Tabel {t_idx} ---")
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            content.append(" | ".join(row_text))

    return "\n".join(content)

@mcp.tool()
def create_or_append_docx(file_path: str, paragraphs: list[str], heading: str = "") -> str:
    """Membuat file .docx baru atau menambahkan paragraf ke file yang sudah ada"""
    try:
        doc = Document(file_path) if os.path.exists(file_path) else Document()

        if heading:
            doc.add_heading(heading, level=1)

        for text in paragraphs:
            doc.add_paragraph(text)

        parent_dir = os.path.dirname(os.path.abspath(file_path))
        if parent_dir:
            os.makedirs(parent_dir, exist_ok=True)

        doc.save(file_path)
        return f"Berhasil menyimpan perubahan ke {file_path}"
    except Exception as e:
        return f"Error: Gagal menyimpan file {file_path}: {e}"

@mcp.tool()
def replace_text_in_docx(file_path: str, old_text: str, new_text: str) -> str:
    """Mencari dan mengganti teks spesifik di dalam paragraf maupun tabel .docx"""
    if not os.path.exists(file_path):
        return f"Error: File {file_path} tidak ditemukan."

    try:
        doc = Document(file_path)
        count = 0

        for p in doc.paragraphs:
            if old_text in p.text:
                count += p.text.count(old_text)
                p.text = p.text.replace(old_text, new_text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if old_text in cell.text:
                        count += cell.text.count(old_text)
                        cell.text = cell.text.replace(old_text, new_text)

        doc.save(file_path)
        return f"Berhasil mengganti {count} kemunculan '{old_text}' menjadi '{new_text}' di {file_path}."
    except Exception as e:
        return f"Error: Gagal memproses file {file_path}: {e}"

if __name__ == "__main__":
    mcp.run()

